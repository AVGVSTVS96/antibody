"""Parsers for CBO cost estimates. Each command prints one JSON document; the TypeScript stage adds citations.

usage: python -I cbo.py xlsx <file.xlsx>      per-title Excel tables (big reconciliation bills) + Summary totals
       python -I cbo.py pdf <file.pdf>        per-section, per-title and per-provision tables in a text PDF
       python -I cbo.py table1 <file.xlsx>    small-bill whole-bill table
       python -I cbo.py glance <file.pdf>     page text, At a Glance box, mandate answers, baseline wording
"""

import itertools
import json
import re
import sys

import openpyxl
import pdfplumber
from openpyxl.cell.rich_text import CellRichText
from openpyxl.utils import get_column_letter


YEAR_COL0 = 3
HIERARCHY = re.compile(r"^(Subtitle|Chapter|Subchapter|Part)\s+([0-9A-Z]+)\.?\s*(.*)$")
TITLE = re.compile(r"^Title\s+([IVXL]+),?\s*(.*)$")


def text(v):
    return re.sub(r"\s+", " ", str(v)).strip() if v is not None else ""


def number(v):
    s = text(v)
    if s in ("", "n.a.", "n.e."):
        return None
    if s == "*":
        return 0
    try:
        return float(s.replace(",", ""))
    except ValueError:
        return None


def measure_kind(label):
    l = label.lower()
    for key, kind in [
        ("budget authority", "budget_authority"),
        ("on-budget revenues", "on_budget_revenues"),
        ("off-budget revenues", "off_budget_revenues"),
        ("on-budget outlays", "on_budget_outlays"),
        ("off-budget outlays", "off_budget_outlays"),
        ("on-budget deficit", "on_budget_deficit"),
        ("off-budget deficit", "off_budget_deficit"),
        ("revenues", "revenues"),
        ("outlays", "outlays"),
        ("deficit", "deficit"),
    ]:
        if key in l:
            return kind
    return "other"


def scorer(label):
    m = re.search(r"\(([^)]*(?:JCT|CBO)[^)]*)\)", label)
    return m.group(1).replace(" Estimate", "") if m else "CBO"


def cell_range(sheet, r0, r1):
    return f"'{sheet}'!A{r0}:O{r1}"


def parse_sheet(ws, years, lines):
    rows = []
    ctx = {"block": None, "Subtitle": None, "Chapter": None, "Subchapter": None, "Part": None}
    title = None
    current = None

    for i, r in enumerate(ws.iter_rows(values_only=True), 1):
        r = list(r) + [None] * 16
        a, b, c = text(r[0]), text(r[1]), text(r[2])
        if not c and b and measure_kind(b) != "other":
            b, c = "", b
        data = r[YEAR_COL0 : YEAR_COL0 + len(years)]
        has_data = any(number(x) is not None for x in data)
        header = text(r[YEAR_COL0])

        if header.startswith("Increases or Decreases (-) in Direct Spending") or header.startswith("Increases in Direct Spending") or header.startswith("Decreases in Direct Spending"):
            ctx["block"], current = "direct_spending", None
            continue
        if "in Revenues" in header and not has_data:
            ctx["block"], current = "revenues", None
            continue
        if header.startswith("Net Increases") or header.startswith("Net Effect"):
            ctx["block"], current = "deficit", None
            continue

        if (m := TITLE.match(a)) and not has_data and title is None:
            title = {"number": m.group(1), "name": m.group(2)}
            continue
        if (m := HIERARCHY.match(a)) and not has_data:
            level = m.group(1)
            ctx[level] = f"{level} {m.group(2)}. {m.group(3)}".strip()
            for lower in ["Subtitle", "Chapter", "Subchapter", "Part"][["Subtitle", "Chapter", "Subchapter", "Part"].index(level) + 1 :]:
                ctx[lower] = None
            current = None
            continue

        if a.startswith("Sec"):
            inline = has_data and measure_kind(c) != "other"
            parent = current if current and current["kind"] == "section" else None
            current = {
                "line": next(lines),
                "kind": "section",
                "section_raw": b,
                "section_title": parent["section_title"] if inline and parent else c,
                "component": None,
                "row0": i,
                "has_rows": False,
                "context": {k: v for k, v in ctx.items() if k != "block"},
            }
            if has_data:
                current["has_rows"] = True
                rows.append(row(ws, title, current, ctx, c if inline else implied_measure(ctx["block"]), data, years, i))
            continue

        if b.lower().startswith("memorand"):
            ctx["block"], current = "memorandum", None
            continue

        if b and not re.match(r"^\d", b) and (not has_data or b.lower().startswith("subtotal")):
            current = {"line": next(lines), "kind": classify_line(b), "section_raw": None, "section_title": b, "component": None, "row0": i, "has_rows": False,
                       "context": {k: v for k, v in ctx.items() if k != "block"}}
            if has_data:
                rows.append(row(ws, title, current, ctx, implied_measure(ctx["block"]) if measure_kind(b) == "other" else b, data, years, i))
            continue

        if a.startswith("Total Changes in") or a.startswith("Net Effect on the Deficit"):
            current = {"line": next(lines), "kind": "total", "section_raw": None, "section_title": a, "component": None, "row0": i, "has_rows": False,
                       "context": {}}
            if has_data:
                rows.append(row(ws, title, current, ctx, a, data, years, i))
            continue

        if a and not has_data and not header:
            current = {"line": next(lines), "kind": classify_line(a), "section_raw": None, "section_title": a, "component": None, "row0": i, "has_rows": False,
                       "context": {k: v for k, v in ctx.items() if k != "block"}}
            continue

        if current and c and not has_data and text(r[YEAR_COL0]):
            current["has_rows"] = True
            rows.append({**row(ws, title, current, ctx, c, [None] * len(years), years, i), "note": text(r[YEAR_COL0]),
                         "included_in": re.findall(r"\d{4,6}", text(r[YEAR_COL0]))})
            continue

        if current and c and not has_data and current["kind"] == "section" and current["has_rows"] and classify_line(c) == "interaction":
            current = {"line": next(lines), "kind": "interaction", "section_raw": None, "section_title": c, "component": None, "row0": i, "has_rows": False,
                       "context": {k: v for k, v in ctx.items() if k != "block"}}
            continue

        if current and c and not has_data:
            if current["has_rows"]:
                current["component"] = c
            elif current["component"]:
                current["component"] = f"{current['component']} {c}"
            else:
                current["section_title"] = f"{current['section_title']} {c}".strip()
            continue

        if current and has_data and c:
            if current["kind"] == "section" and b:
                current["section_raw"] = f"{current['section_raw']} {b}"
            current["has_rows"] = True
            rows.append(row(ws, title, current, ctx, c, data, years, i))
    for r in rows:
        line = r.pop("_line")
        r["line"] = line["line"]
        r["section_raw"] = line["section_raw"]
        r["sections"] = expand_sections(line["section_raw"]) if line["section_raw"] else []
        r["range"] = cell_range(ws.title, line["row0"], int(r["range"].split(":O")[1]))
    return title, rows


def implied_measure(block):
    return {"revenues": "Estimated Revenues", "deficit": "Effect on the Deficit"}.get(block, "Estimated Outlays")


def row(ws, title, current, ctx, label, data, years, i):
    return {
        "sheet": ws.title,
        "title": title,
        **current["context"],
        "block": ctx["block"],
        "line_kind": current["kind"],
        "_line": current,
        "section_title": current["section_title"],
        "component": current["component"],
        "measure": label,
        "measure_kind": measure_kind(label),
        "scored_by": scorer(label),
        "values": dict(zip(map(str, years), [number(x) for x in data])),
        "raw": dict(zip(map(str, years), [text(x) for x in data])),
        "range": cell_range(ws.title, current["row0"], i),
    }


def classify_line(b):
    l = b.lower()
    if "interaction" in l:
        return "interaction"
    if l.startswith("subtotal") or re.search(r"\btotal\b", l):
        return "subtotal"
    if "sequestration" in l:
        return "sequestration"
    return "other"


def expand_sections(raw):
    out = []
    for part in re.split(r"\s*(?:,|and|;)\s*", raw.strip()):
        if m := re.fullmatch(r"(\d+)-(\d+)", part):
            lo, hi = int(m.group(1)), int(m.group(2))
            if 0 <= hi - lo < 200:
                out += [str(n) for n in range(lo, hi + 1)]
                continue
        out += re.findall(r"\d+(?:\([a-z0-9]+\)(?:-\([a-z0-9]+\))?)?", part)
    return out or [raw]


def find_years(ws):
    """Column labels of a sheet. Some CBO sheets print the two total columns over two rows (' 2021-' above '2030'),
    which must read as '2021-2030', not collide with the annual 2030 column."""
    rows = list(ws.iter_rows(min_row=1, max_row=30, values_only=True))
    for k, r in enumerate(rows):
        if len(r) > YEAR_COL0 and text(r[YEAR_COL0]).isdigit():
            above = rows[k - 1] if k else ()
            out = []
            for j in range(YEAR_COL0, len(r)):
                if r[j] is None:
                    continue
                label = text(r[j]).replace(" ", "")
                lead = text(above[j]).replace(" ", "") if j < len(above) and above[j] is not None else ""
                out.append(lead + label if lead.endswith("-") and label.isdigit() else label)
            return out
    return None


def summary_totals(ws, years):
    """The 'Total Changes' block of a Summary sheet: whole-bill outlays, revenues and deficit, as CBO prints them."""
    rows, inside = [], False
    for i, r in enumerate(ws.iter_rows(values_only=True), 1):
        r = list(r) + [None] * 16
        a, b = text(r[0]), text(r[1])
        if a.startswith("Total Changes"):
            inside = True
            continue
        if not inside:
            continue
        data = r[YEAR_COL0 : YEAR_COL0 + len(years)]
        if not b or not any(number(v) is not None for v in data):
            break
        rows.append({"measure": b, "measure_kind": measure_kind(b), "values": dict(zip(years, [number(v) for v in data])),
                     "raw": dict(zip(years, [text(v) for v in data])), "range": f"'{ws.title}'!A{i}:O{i}"})
    return rows


CONNECTIVE = re.compile(r"(?:,|\b(?:and|of|for|the|to|in|on|or|with))$", re.I)
PROVISION_BLOCKS = [("Direct Spending", "direct_spending"), ("Discretionary", "discretionary"), ("Revenues", "revenues"), ("Deficit", "deficit")]


def plain(v):
    """A cell's text without its superscript runs, which are footnote marks."""
    if isinstance(v, CellRichText):
        v = "".join(p if isinstance(p, str) else ("" if p.font.vertAlign == "superscript" else p.text) for p in v)
    return text(v)


def is_provision_table(ws):
    """A numbered table that lists provisions by name (no section numbers) under 'in Direct Spending' headings."""
    return ws.title.startswith("Table") and any(
        "in Direct Spending" in text(r[YEAR_COL0]) for r in ws.iter_rows(max_row=40, values_only=True) if len(r) > YEAR_COL0
    )


def provision_sheet(ws, years, lines):
    """Rows of a table that names provisions instead of sections, with the heading each sits under and its parent
    heading. Amounts are returned in millions, however the sheet prints them (FRA's tables are in billions)."""
    n = len(years)
    billions = any("Billions" in text(c) for r in ws.iter_rows(max_row=14, values_only=True) for c in r)
    scale = 1000 if billions else 1
    last = get_column_letter(YEAR_COL0 + n)
    rows, block, heads, current = [], None, [], None

    def start(kind, title, parent, row):
        return {"line": next(lines), "kind": kind, "title": title, "parent": parent, "row0": row, "has_rows": False}

    def add(cur, label, data, i):
        cur["has_rows"] = True
        if block == "discretionary":
            return
        nums = [number(x) for x in data]
        rows.append({
            "sheet": ws.title, "block": block, "line_kind": cur["kind"], "line": cur["line"], "sections": [],
            "section_title": cur["title"], "parent": cur["parent"], "component": None, "measure": label,
            "measure_kind": measure_kind(label), "scored_by": "CBO",
            "values": dict(zip(years, [None if x is None else round(x * scale, 6) for x in nums])),
            "raw": dict(zip(years, [text(x) for x in data])),
            "range": f"'{ws.title}'!A{cur['row0']}:{last}{i}",
        })

    def under_heads():
        (title, row), parent = heads[-1], heads[-2][0] if len(heads) > 1 else None
        return start("program", title, parent, row)

    for i, r in enumerate(ws.iter_rows(values_only=True), 1):
        r = list(r) + [None] * (YEAR_COL0 + n + 1)
        a, b, c = plain(r[0]), plain(r[1]), plain(r[2])
        data = r[YEAR_COL0 : YEAR_COL0 + n]
        has = any(number(x) is not None for x in data)
        if a.startswith("Source"):
            break
        if [text(x) for x in data] == years:
            continue
        if not has and (head := plain(r[YEAR_COL0])):
            for key, name in PROVISION_BLOCKS:
                if key in head:
                    block, heads, current = name, [], None
        elif has and (b or c):
            if current is None:
                current = under_heads() if heads else start("total", b or c, None, i)
            add(current, b or c, data, i)
        elif has:
            if heads and CONNECTIVE.search(heads[-1][0]):
                title, row = heads.pop()
                current = start("program", f"{title} {a}", heads[-1][0] if heads else None, row)
            else:
                current = start("program", a, heads[-1][0] if heads else None, i)
            add(current, implied_measure(block), data, i)
            heads = []
        elif a and not b and not c:
            if heads and CONNECTIVE.search(heads[-1][0]):
                heads[-1] = (f"{heads[-1][0]} {a}", heads[-1][1])
            elif current and current["has_rows"]:
                heads = [(a, i)]
            else:
                heads.append((a, i))
            current = None
        elif label := b or c:
            kind = "total" if label.lower().startswith("total") else "subtotal" if label.lower().startswith("subtotal") else None
            if kind:
                current, heads = start(kind, label, None, i), []
    return rows, ("billions" if billions else None)


def provision_totals(rows):
    """Whole-bill direct spending and deficit as a provision table totals them."""
    return [
        {"block": r["block"], "measure": r["measure"], "measure_kind": r["measure_kind"], "values": r["values"], "range": r["range"]}
        for r in rows
        if r["line_kind"] == "total" and r["measure_kind"] in ("budget_authority", "outlays", "deficit")
    ]


def xlsx(path):
    wb = openpyxl.load_workbook(path, data_only=True, rich_text=True)
    lines = itertools.count(1)
    out = {"format": "xlsx", "sheets": [], "summary": None}
    for ws in wb.worksheets:
        years = find_years(ws)
        if not years:
            continue
        if ws.title.startswith("Summary"):
            out["summary"] = {"sheet": ws.title, "years": years, "rows": summary_totals(ws, years)}
        if is_provision_table(ws):
            rows, unit = provision_sheet(ws, years, lines)
            out["sheets"].append({"sheet": ws.title, "title": None, "years": years, "unit": unit, "rows": rows})
            if not out["summary"]:
                out["summary"] = {"sheet": ws.title, "years": years, "rows": provision_totals(rows)}
            continue
        if not ws.title.startswith("Title"):
            continue
        title, rows = parse_sheet(ws, years, lines)
        out["sheets"].append({"sheet": ws.title, "title": title, "years": years, "rows": rows})
    return out
NUM = re.compile(r"^(-?\d{1,3}(,\d{3})*|-?\d+|\*|n\.e\.|n\.a\.)$")
YEARS = re.compile(r"^(\d{4}\s+){9,}\d{4}-\d{4}\s+\d{4}-\d{4}$")
HIERARCHY = re.compile(r"^(Subtitle|Chapter|Subchapter|Part)\s+([0-9A-Z]+)\.\s*(.*)$")
SEC_NUMBER = r"\d+[A-Za-z]?(?:\([a-z0-9]+\))*"
SEC = re.compile(r"^Secs?\.\s+(\d[\d\-–(),a-z ]*?)\s*(?:\s([A-Z].*))?$")
SEC_WORD = re.compile(rf"^Sections?\s+({SEC_NUMBER}(?:\s*(?:,\s*and|,|and|–|-)\s*{SEC_NUMBER})*)\.(?:\s+([A-Z].*))?$")
TABLE_HEAD = re.compile(r"^Table (\d+)(?:\.\s*(.*)|)$")
TABLE_TITLE = re.compile(r"Estimated Budgetary Effects of (Title [IVXL]+), ([^,]+(?:, [^,]+)*?), of Public Law")
TABLE_DIVISION = re.compile(r"\bEffects of Division ([A-Z]{1,2})\b")
DIVISION = re.compile(r"^(?:DIVISION|Division)\s+([A-Z]{1,2})\s*[—–.:-]+\s*(.*)$")
UNIT_TITLE = re.compile(r"^(?:TITLE|Title)\s+([IVXLC]+)\s*[—–.:-]+\s*(.*)$")
UNIT_TOTAL = re.compile(r"^(?:(?:DIVISION|Division)\s+[A-Z]{1,2}|(?:TITLE|Title)\s+[IVXLC]+)\s+Total$|^Total,\s*(?:Division|Title)\s+\S+$")
ON_OFF_BUDGET = re.compile(r"^(On|Off)-budget(?:\s+(?:revenues|outlays))?$", re.I)
END_OF_TABLE = re.compile(r"^(?:Sources?:|Estimates are relative|Components may not sum|Notes$)")
LEAF = ("section", "title", "program")
STRUCTURE = ("division", "title", "program")
SUPERSCRIPT_MAX = 7.0


def value(tok):
    if tok == "*":
        return 0.0
    if tok in ("n.e.", "n.a."):
        return None
    return float(tok.replace(",", ""))


def split_numbers(line, n):
    toks = line.split()
    tail = []
    while toks and NUM.match(toks[-1]) and len(tail) < n:
        tail.insert(0, toks.pop())
    if len(tail) == n:
        return " ".join(toks), tail
    return line, None


def parse_years(line, above):
    """Column labels of a PDF table. Some tables print the two total columns as '2021- 2021-' above a row of plain
    years ('... 2031 2026 2031'), which must read as '2021-2026' and '2021-2031'."""
    if YEARS.match(line):
        return line.split()
    toks = line.split()
    if len(toks) >= 12 and all(re.fullmatch(r"\d{4}", t) for t in toks):
        annual = toks[:-2]
        lead = re.findall(r"(\d{4})-(?=\s|$)", above)
        if len(lead) >= 2 and all(int(b) == int(a) + 1 for a, b in zip(annual, annual[1:])):
            return annual + [f"{lead[-2]}-{toks[-2]}", f"{lead[-1]}-{toks[-1]}"]
    return None


def section_label(label):
    """(section numbers as CBO printed them, the title that follows) for 'Sec. 124 Title' and 'Section 124. Title'."""
    if m := SEC.match(label):
        return m.group(1).strip(), m.group(2) or ""
    if m := SEC_WORD.match(label):
        return m.group(1).strip(), m.group(2) or ""
    return None


def page_lines(page):
    """A page's text lines with their left edge; small superscript characters (footnote marks) are left out."""
    clean = page.filter(lambda o: o.get("object_type") != "char" or o["size"] > SUPERSCRIPT_MAX)
    return [(l["text"].strip(), l["x0"]) for l in clean.extract_text_lines()]


def per_section_pdf(pdf_path):
    with pdfplumber.open(pdf_path) as pdf:
        pages = [page_lines(page) for page in pdf.pages]
    return {"format": "pdf", "sheets": parse_tables(pages), "summary": pdf_summary(pdf_path)}


def parse_tables(pages):
    """Per-section, per-title and per-provision rows of the numbered tables in a PDF, from its pages' text lines
    (`(text, left edge)` pairs)."""
    lines_seq = itertools.count(1)
    rows, years = [], None
    table = title = block = None
    skip = reading = False
    caption, heading, above = [], set(), ""
    unit = {"division": None, "title": None}
    ctx = {"Subtitle": None, "Chapter": None, "Subchapter": None, "Part": None}
    current = None
    for pno, lines in enumerate(pages, 1):
        if any(t.startswith("Table") and "Notes" in t for t, _ in lines[:6]):
            continue
        for line, x0 in lines:
            if not line or line.isdigit() or line.startswith(("See also", "How CBO", "Congressional Budget Office")) or line == "Cost Estimate":
                continue
            before, above = above, line
            if m := TABLE_HEAD.match(line):
                if table != m.group(1):
                    table, title, block, current, years, skip = m.group(1), None, None, None, None, False
                    ctx, unit = dict.fromkeys(ctx), {"division": None, "title": None}
                    heading = set()
                caption = [m.group(2)] if m.group(2) else []
                heading.update(caption)
                reading = years is not None
                continue
            if not table:
                continue
            if found := parse_years(line, before):
                years, reading = found, True
                head = " ".join(caption)
                skip = skip or "Discretionary Spending" in head or head.startswith("Summary")
                if (m := TABLE_DIVISION.search(head)) and not unit["division"]:
                    unit = {"division": m.group(1), "title": None}
                continue
            if not reading:
                caption.append(line)
                heading.add(line)
                if m := TABLE_TITLE.search(line):
                    title = {"number": m.group(1).split()[1], "name": m.group(2)}
                continue
            if line in heading or line.startswith("By Fiscal Year"):
                continue
            if skip or not years:
                continue
            if END_OF_TABLE.match(line):
                years, current = None, None
                continue
            if "in Direct Spending" in line and line.startswith(("Increases", "Decreases")):
                block, current = "direct_spending", None
                continue
            if "in Revenues" in line and line.startswith(("Increases", "Decreases")):
                block, current = "revenues", None
                continue
            if "Discretionary Spending" in line and line.startswith(("Increases", "Decreases")):
                block, current = "discretionary", None
                continue
            if line.startswith("Net Increase") or line.startswith("From Changes"):
                block, current = "deficit", None
                continue
            if line.startswith("Memorandum"):
                block, current = "memorandum", None
                continue
            if block == "discretionary":
                continue
            if m := HIERARCHY.match(line):
                level = m.group(1)
                ctx[level] = f"{level} {m.group(2)}. {m.group(3)}"
                for lower in list(ctx)[list(ctx).index(level) + 1 :]:
                    ctx[lower] = None
                current = None
                continue

            label, nums = split_numbers(line, len(years))
            note = None
            if " Estimate included in" in label:
                label, note = label.split(" Estimate included in", 1)
                note = "Estimate included in" + note

            if (m := ON_OFF_BUDGET.match(label)) and current and nums:
                label = f"{m.group(1).title()}-Budget {'Revenues' if block == 'revenues' else 'Outlays'}"
            sec = section_label(label)
            if sec and not (current and current.get("joining")):
                current = new_line(lines_seq, "section", sec[1], pno, ctx, sec[0], unit, x0)
            elif sec and current and current.get("joining"):
                current["section_raw"] += f" and {sec[0]}"
                current["section_title"] += f" and {sec[1]}"
                current["joining"] = False
            elif label == "and" and current and current["kind"] == "section":
                current["joining"] = True
                continue
            elif label.startswith(("Total Changes in", "Net Effect on the Deficit")):
                current = new_line(lines_seq, "total", label, pno, {}, None, unit, x0)
            elif UNIT_TOTAL.match(label) or label.startswith(("Subtotal", "Interaction", "Total Interaction")):
                current = new_line(lines_seq, classify_line(label), label, pno, ctx, None, unit, x0)
            elif m := DIVISION.match(label):
                unit = {"division": m.group(1), "title": None}
                current = new_line(lines_seq, "division", m.group(2), pno, ctx, None, unit, x0)
            elif m := UNIT_TITLE.match(label):
                unit = {**unit, "title": m.group(1)}
                current = new_line(lines_seq, "title", m.group(2), pno, ctx, None, unit, x0)
            elif current and not nums and not note:
                opens = current["kind"] == "program" or (current["kind"] in ("division", "title") and not current["has_rows"])
                if (unit["division"] or unit["title"]) and opens:
                    wraps = not current["has_rows"] and (current["kind"] == "program" or x0 < current["x0"] + 3)
                    if wraps:
                        current["section_title"] += f" {label}"
                    else:
                        current = new_line(lines_seq, "program", label, pno, ctx, None, unit, x0)
                elif current["has_rows"]:
                    current["component"] = label
                else:
                    current["section_title"] += f" {label}"
                continue
            elif current and nums:
                if measure_kind(label) == "other" and current["kind"] in LEAF and not current["has_rows"]:
                    current["section_title"] += f" {label}"
                    label = implied_measure(block)
                rows.append(emit(table, title, block, current, label, years, nums, pno, lines_seq))
                continue
            else:
                continue

            if note:
                current["has_rows"] = True
                rows.append({**emit(table, title, block, current, "", years, None, pno, lines_seq), "note": note,
                             "included_in": re.findall(r"\d{4,6}", note)})
            elif nums:
                rows.append(emit(table, title, block, current, implied_measure(block), years, nums, pno, lines_seq))
    return group(rows)


def pdf_summary(pdf_path):
    """Whole-bill totals printed in a PDF's tables, before any per-title 'Table N.': total direct spending,
    total discretionary spending, total revenues and the effect on the deficit."""
    rows, years, after, memo = [], None, "", None
    with pdfplumber.open(pdf_path) as pdf:
        for pno, page in enumerate(pdf.pages, 1):
            clean = page.filter(lambda o: o.get("object_type") != "char" or o["size"] > SUPERSCRIPT_MAX)
            lines = [l.strip() for l in (clean.extract_text() or "").splitlines()]
            for line in lines:
                if rows and re.match(r"^Table \d+\.?(\s|$)", line):
                    return {"sheet": "Summary", "years": years, "rows": rows}
                if YEARS.match(line):
                    years = line.split()
                    continue
                if not years:
                    continue
                label, nums = split_numbers(line, len(years))
                if label.startswith("Memorandum"):
                    memo = ""
                    continue
                if memo is not None:
                    if not nums:
                        memo = f"{memo} {label}".strip()
                        continue
                    if memo.startswith("Nonscorable"):
                        rows.append({"block": "memorandum", "measure": f"{memo} {label}".strip(),
                                     "measure_kind": "nonscorable_revenues",
                                     "values": dict(zip(years, [value(t) for t in nums])), "raw": dict(zip(years, nums)),
                                     "range": f"page {pno}"})
                    memo = None
                    continue
                if label.startswith(("Total Changes in", "Net Increase or Decrease", "Increases or Decreases", "Increases in", "Decreases in")):
                    after = label
                if not nums:
                    continue
                direct = after.startswith("Total Changes in Direct Spending")
                discretionary = after.startswith("Total Changes in") and ("Discretionary" in after or "Subject to Appropriation" in after)
                if label.startswith("Total Changes in Revenues"):
                    block = "revenues"
                elif (direct or discretionary) and label in ("Budget Authority", "Estimated Budget Authority", "Estimated Outlays"):
                    block = "discretionary" if discretionary else "direct_spending"
                elif after.startswith("Net Increase or Decrease") and label == "Effect on the Deficit":
                    block = "deficit"
                else:
                    continue
                if not any(r["block"] == block and r["measure"] == label for r in rows):
                    rows.append({"block": block, "measure": label, "measure_kind": measure_kind(label),
                                 "values": dict(zip(years, [value(t) for t in nums])), "raw": dict(zip(years, nums)),
                                 "range": f"page {pno}"})
    return {"sheet": "Summary", "years": years, "rows": rows} if rows else None


def new_line(seq, kind, text, pno, ctx, section_raw, unit=None, x0=0.0):
    return {"line": None if kind in STRUCTURE else next(seq), "kind": kind, "section_raw": section_raw, "section_title": text, "component": None, "page0": pno,
            "has_rows": False, "joining": False, "context": dict(ctx), "unit": dict(unit or {}), "x0": x0}


def emit(table, title, block, current, measure, years, nums, pno, seq):
    current["has_rows"] = True
    if current["line"] is None:
        current["line"] = next(seq) if current["kind"] != "division" else 0
    pages = sorted({current["page0"], pno})
    return {
        "table": table,
        "sheet": f"Table {table}",
        "title": title,
        **current["context"],
        "division": current["unit"].get("division"),
        "title_no": current["unit"].get("title"),
        "block": block,
        "line_kind": current["kind"],
        "_line": current,
        "section_title": current["section_title"],
        "component": current["component"],
        "measure": measure,
        "measure_kind": measure_kind(measure) if measure else "note",
        "scored_by": scorer(measure),
        "values": dict(zip(years, [value(t) for t in nums])) if nums else dict.fromkeys(years),
        "raw": dict(zip(years, nums)) if nums else {},
        "pages": pages,
        "range": f"page {pages[0]}" if len(pages) == 1 else f"pages {pages[0]}-{pages[-1]}",
    }


def group(rows):
    for r in rows:
        line = r.pop("_line")
        r["line"] = line["line"]
        r["section_raw"] = line["section_raw"]
        r["sections"] = expand_sections(line["section_raw"]) if line["section_raw"] else []
        r["_line_id"] = id(line)
    with_totals = {(r["_line_id"], r["block"]) for r in rows if ", Total" in r["measure"]}
    tables = {}
    for r in rows:
        r["breakdown_of_total"] = bool(r["component"]) and (r.pop("_line_id"), r["block"]) in with_totals
        t = tables.setdefault(r["table"], {"sheet": r["sheet"], "title": r["title"], "years": list(r["values"]), "rows": []})
        t["rows"].append(r)
    return list(tables.values())


def table1(path):
    ws = openpyxl.load_workbook(path, data_only=True).active
    rows, years, block = [], None, None
    for i, r in enumerate(ws.iter_rows(values_only=True), 1):
        label, cells = text(r[0]), list(r[1:])
        if not label and text(cells[0]).isdigit():
            years = [text(x) for x in cells if x is not None]
            continue
        head = text(cells[0]) if not label else ""
        if head.startswith("From Changes"):
            pass
        elif "Deficit" in head:
            block = "deficit"
        elif "Subject to Appropriation" in head:
            block = "discretionary"
        elif "Direct Spending" in head:
            block = "direct_spending"
        elif "Revenues" in head:
            block = "revenues"
        if label and years and any(number(x) is not None for x in cells[: len(years)]):
            vals = cells[: len(years)]
            rows.append({
                "sheet": ws.title, "block": block, "line_kind": "whole_bill", "sections": [], "section_title": None,
                "measure": label, "measure_kind": measure_kind(label), "scored_by": "CBO",
                "values": dict(zip(years, [number(x) for x in vals])), "raw": dict(zip(years, [text(x) for x in vals])),
                "range": f"'{ws.title}'!A{i}:{chr(65 + len(years))}{i}",
            })
    return {"format": "table1", "sheets": [], "summary": {"sheet": ws.title, "years": years, "rows": rows}}
GLANCE_VALUE = re.compile(r"^(-?[\d,]+(?:\.\d+)?|\*|n\.e\.|n\.a\.|-)$")
PERIOD = re.compile(r"^\d{4}(?:-\d{4})?$")
MONTH_BASELINE = re.compile(r"relative to CBO[’']s ((?:January|February|March|April|May|June|July|August|September|October|November|December) \d{4}) baseline", re.I)


def glance_value(tok):
    if tok == "*":
        return 0.0
    if tok in ("n.e.", "n.a.", "not estimated", "-"):
        return None
    return float(tok.replace(",", ""))


def word_lines(words, tolerance=3):
    lines = []
    for w in sorted(words, key=lambda w: (w["top"], w["x0"])):
        if lines and abs(lines[-1][0]["top"] - w["top"]) <= tolerance:
            lines[-1].append(w)
        else:
            lines.append([w])
    return [sorted(l, key=lambda w: w["x0"]) for l in lines]


def mid(ws):
    return sum((w["top"] + w["bottom"]) / 2 for w in ws) / len(ws)


def glance_table(words):
    """The 'At a Glance' box on page 1: whole-bill direct spending, revenues, deficit and spending subject to
    appropriation. Read by word position, because labels wrap onto two lines around their numbers."""
    lines = word_lines(words)
    head = next((l for l in lines if "Millions" in " ".join(w["text"] for w in l)), None)
    if head is None:
        return None
    periods = [w for w in head if PERIOD.match(w["text"])]
    if not periods:
        return None
    left = min(w["x0"] for w in periods) - 8
    end = next((l[0]["top"] for l in lines if l[0]["top"] > head[0]["top"]
                and " ".join(w["text"] for w in l).startswith(("Increases net", "Statutory", "* =", "Mandate"))), float("inf"))
    body = [l for l in lines if head[0]["bottom"] < l[0]["top"] < end]
    value_rows, label_lines = [], []
    for l in body:
        labels = [w for w in l if w["x1"] < left]
        cells = [w for w in l if w["x1"] >= left]
        text = " ".join(w["text"] for w in cells).replace("not estimated", "n.e.")
        toks = text.split()
        if cells and len(toks) == len(periods) and all(GLANCE_VALUE.match(t) for t in toks):
            value_rows.append({"mid": mid(cells), "values": toks, "label": labels})
        elif labels:
            label_lines.append(labels)
    for l in label_lines:
        nearest = min(value_rows, key=lambda r: abs(r["mid"] - mid(l)), default=None)
        if nearest and abs(nearest["mid"] - mid(l)) < 16:
            nearest["label"] = nearest["label"] + l
    rows = []
    for r in value_rows:
        label = " ".join(w["text"] for w in sorted(r["label"], key=lambda w: (round(w["top"]), w["x0"])))
        rows.append({"label": label, "values": dict(zip([p["text"] for p in periods], [glance_value(v) for v in r["values"]])),
                     "raw": dict(zip([p["text"] for p in periods], r["values"]))})
    return {"periods": [p["text"] for p in periods], "rows": rows} if rows else None


def mandates(words):
    lines = word_lines(words)
    questions = {}
    for l in lines:
        text = " ".join(w["text"] for w in l)
        for key, q in (("intergovernmental", "intergovernmental mandate?"), ("private_sector", "private-sector mandate?")):
            if q in text:
                end = next(w for w in l if w["text"].endswith("mandate?"))
                questions[key] = (mid([end]), end["x1"])
    if not questions:
        return {"intergovernmental": None, "private_sector": None}
    right = max(x for _, x in questions.values()) + 4
    answers = {k: [] for k in questions}
    for l in lines:
        cells = [w for w in l if w["x0"] > right]
        if not cells:
            continue
        key = min(questions, key=lambda k: abs(questions[k][0] - mid(cells)))
        if abs(questions[key][0] - mid(cells)) < 16:
            answers[key].append(" ".join(w["text"] for w in cells))
    out = {k: " ".join(v) or None for k, v in answers.items()}
    return {"intergovernmental": out.get("intergovernmental"), "private_sector": out.get("private_sector")}


def glance(path):
    with pdfplumber.open(path) as pdf:
        pages = [re.sub(r"[ \t]+", " ", p.extract_text() or "").strip() for p in pdf.pages]
        words = pdf.pages[0].extract_words() if pdf.pages else []
    flat = re.sub(r"\s+", " ", " ".join(pages))
    baseline = MONTH_BASELINE.search(flat)
    return {
        "format": "glance",
        "pages": pages,
        "glance": glance_table(words),
        "mandates": mandates(words),
        "baseline_month": baseline.group(1) if baseline else None,
        "budget_enforcement_baseline": bool(re.search(r"budget[- ]enforcement baseline|baseline used for budget enforcement", flat, re.I)),
    }


def pages(path):
    """The `pdf` tables from a JSON list of pages of `[text, left edge]` lines, for tests that have no PDF."""
    with open(path) as f:
        return {"format": "pdf", "sheets": parse_tables([[(t, x0) for t, x0 in page] for page in json.load(f)]), "summary": None}


if __name__ == "__main__":
    command, path = sys.argv[1], sys.argv[2]
    parse = {"xlsx": xlsx, "pdf": per_section_pdf, "pages": pages, "table1": table1, "glance": glance}[command]
    json.dump(parse(path), sys.stdout)
