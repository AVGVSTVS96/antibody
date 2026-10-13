import json, random, re, sys
import openpyxl

rows = json.load(open(sys.argv[1]))
files = {"hr1-pl119-21-clb": "61570-pl119-21-2025Recon-CLB.xlsx", "hr1-pl119-21-beb": "61569-pl119-21-2025Recon-BEB.xlsx", "hr1-house-passed": "HR1_HousePassed_6-4-2025.xlsx"}
wbs = {k: openpyxl.load_workbook(sys.argv[2] + "/" + v, data_only=True) for k, v in files.items()}
random.seed(int(sys.argv[3]))
cand = [r for r in rows if r["estimate"] in files and r["line_kind"] == "section"]
ok = 0
sample = random.sample(cand, int(sys.argv[4]))
for r in sample:
    sh, a, b = re.search(r"\('([^']+)'!A(\d+):O(\d+)\)", r["source"]["cite"]).groups()
    ws = wbs[r["estimate"]][sh]
    vals = [c.value for c in ws[int(b)] if c.value is not None]
    head = " ".join(str(c.value) for c in ws[int(a)] if c.value is not None)
    want = r["ten_year_total"]
    last = vals[-1]
    val_ok = last == want or (last == "*" and abs(want or 0) < 1) or (isinstance(last, str) and want is None)
    good = val_ok and r["sections"][0] in head and str(vals[0]).strip() == r["measure"].strip()
    ok += good
    if not good:
        print("MISMATCH", r["estimate"], r["sections"], r["measure"], want, "|", head[:50], vals[0], last)
print(f"{ok}/{len(sample)}")
