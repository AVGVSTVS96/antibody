import json
import re
import sys

import openpyxl

YEAR_COL0 = 3
HIERARCHY = re.compile(r"^(Subtitle|Chapter|Subchapter|Part)\s+([0-9A-Z]+)\.?\s*(.*)$")
TITLE = re.compile(r"^Title\s+([IVXL]+),?\s*(.*)$")


def text(v):
    return re.sub(r"\s+", " ", str(v)).strip() if v is not None else ""


def number(v):
    s = text(v)
    if s in ("", "n.a.", "n.e."):
        return None
    if s == "*":
        return 0
    try:
        return float(s.replace(",", ""))
    except ValueError:
        return None


def measure_kind(label):
    l = label.lower()
    for key, kind in [
        ("budget authority", "budget_authority"),
        ("on-budget revenues", "on_budget_revenues"),
        ("off-budget revenues", "off_budget_revenues"),
        ("on-budget outlays", "on_budget_outlays"),
        ("off-budget outlays", "off_budget_outlays"),
        ("on-budget deficit", "on_budget_deficit"),
        ("off-budget deficit", "off_budget_deficit"),
        ("revenues", "revenues"),
        ("outlays", "outlays"),
        ("deficit", "deficit"),
    ]:
        if key in l:
            return kind
    return "other"


def scorer(label):
    m = re.search(r"\(([^)]*(?:JCT|CBO)[^)]*)\)", label)
    return m.group(1).replace(" Estimate", "") if m else "CBO"


def cell_range(sheet, r0, r1):
    return f"'{sheet}'!A{r0}:O{r1}"


def parse_sheet(ws, years):
    rows = []
    ctx = {"block": None, "Subtitle": None, "Chapter": None, "Subchapter": None, "Part": None}
    title = None
    current = None

    for i, r in enumerate(ws.iter_rows(values_only=True), 1):
        r = list(r) + [None] * 16
        a, b, c = text(r[0]), text(r[1]), text(r[2])
        data = r[YEAR_COL0 : YEAR_COL0 + len(years)]
        has_data = any(number(x) is not None for x in data)
        header = text(r[YEAR_COL0])

        if header.startswith("Increases or Decreases (-) in Direct Spending") or header.startswith("Increases in Direct Spending") or header.startswith("Decreases in Direct Spending"):
            ctx["block"], current = "direct_spending", None
            continue
        if "in Revenues" in header and not has_data:
            ctx["block"], current = "revenues", None
            continue
        if header.startswith("Net Increases") or header.startswith("Net Effect"):
            ctx["block"], current = "deficit", None
            continue

        if (m := TITLE.match(a)) and not has_data and title is None:
            title = {"number": m.group(1), "name": m.group(2)}
            continue
        if (m := HIERARCHY.match(a)) and not has_data:
            level = m.group(1)
            ctx[level] = f"{level} {m.group(2)}. {m.group(3)}".strip()
            for lower in ["Subtitle", "Chapter", "Subchapter", "Part"][["Subtitle", "Chapter", "Subchapter", "Part"].index(level) + 1 :]:
                ctx[lower] = None
            current = None
            continue

        if a.startswith("Sec"):
            inline = has_data and measure_kind(c) != "other"
            parent = current if current and current["kind"] == "section" else None
            current = {
                "kind": "section",
                "section_raw": b,
                "section_title": parent["section_title"] if inline and parent else c,
                "component": None,
                "row0": i,
                "has_rows": False,
                "context": {k: v for k, v in ctx.items() if k != "block"},
            }
            if inline:
                current["has_rows"] = True
                rows.append(row(ws, title, current, ctx, c, data, years, i))
            continue

        if b and not has_data and not re.match(r"^\d", b):
            current = {"kind": classify_line(b), "section_raw": None, "section_title": b, "component": None, "row0": i, "has_rows": False,
                       "context": {k: v for k, v in ctx.items() if k != "block"}}
            continue

        if a.startswith("Total Changes in") or a.startswith("Net Effect on the Deficit"):
            current = {"kind": "total", "section_raw": None, "section_title": a, "component": None, "row0": i, "has_rows": False,
                       "context": {}}
            if has_data:
                rows.append(row(ws, title, current, ctx, a, data, years, i))
            continue

        if current and c and not has_data and text(r[YEAR_COL0]):
            current["has_rows"] = True
            rows.append({**row(ws, title, current, ctx, c, [None] * len(years), years, i), "note": text(r[YEAR_COL0]),
                         "included_in": re.findall(r"\d{4,6}", text(r[YEAR_COL0]))})
            continue

        if current and c and not has_data:
            if current["has_rows"]:
                current["component"] = c
            elif current["component"]:
                current["component"] = f"{current['component']} {c}"
            else:
                current["section_title"] = f"{current['section_title']} {c}".strip()
            continue

        if current and has_data and c:
            if current["kind"] == "section" and b:
                current["section_raw"] = f"{current['section_raw']} {b}"
            current["has_rows"] = True
            rows.append(row(ws, title, current, ctx, c, data, years, i))
    for r in rows:
        line = r.pop("_line")
        r["section_raw"] = line["section_raw"]
        r["sections"] = expand_sections(line["section_raw"]) if line["section_raw"] else []
        r["range"] = cell_range(ws.title, line["row0"], int(r["range"].split(":O")[1]))
    return title, rows


def row(ws, title, current, ctx, label, data, years, i):
    return {
        "sheet": ws.title,
        "title": title,
        **current["context"],
        "block": ctx["block"],
        "line_kind": current["kind"],
        "_line": current,
        "section_title": current["section_title"],
        "component": current["component"],
        "measure": label,
        "measure_kind": measure_kind(label),
        "scored_by": scorer(label),
        "values": dict(zip(map(str, years), [number(x) for x in data])),
        "raw": dict(zip(map(str, years), [text(x) for x in data])),
        "range": cell_range(ws.title, current["row0"], i),
    }


def classify_line(b):
    l = b.lower()
    if "interaction" in l:
        return "interaction"
    if l.startswith("subtotal") or l.startswith("total"):
        return "subtotal"
    if "sequestration" in l:
        return "sequestration"
    return "other"


def expand_sections(raw):
    out = []
    for part in re.split(r"\s*(?:,|and|;)\s*", raw.strip()):
        if m := re.fullmatch(r"(\d+)-(\d+)", part):
            lo, hi = int(m.group(1)), int(m.group(2))
            if 0 <= hi - lo < 200:
                out += [str(n) for n in range(lo, hi + 1)]
                continue
        out += re.findall(r"\d+(?:\([a-z0-9]+\)(?:-\([a-z0-9]+\))?)?", part)
    return out or [raw]


def find_years(ws):
    for r in ws.iter_rows(min_row=1, max_row=30, values_only=True):
        if text(r[YEAR_COL0]).isdigit():
            return [text(x).replace(" ", "") for x in r[YEAR_COL0:] if x is not None]
    raise ValueError(f"no year header in {ws.title}")


def main(path, meta_path):
    meta = json.load(open(meta_path))
    wb = openpyxl.load_workbook(path, data_only=True)
    out = {"meta": meta, "sheets": []}
    for ws in wb.worksheets:
        if not ws.title.startswith("Title"):
            continue
        years = find_years(ws)
        title, rows = parse_sheet(ws, years)
        out["sheets"].append({"sheet": ws.title, "title": title, "years": years, "rows": rows})
    json.dump(out, sys.stdout, indent=1)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
